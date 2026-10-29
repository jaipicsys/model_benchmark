#!/usr/bin/env python3
"""
VideoMAE long-video detection logger -> Excel (with resume cache).

Runs sliding-window VideoMAE inference on one video and writes an Excel file
with exactly three columns:
    Detection | Video Timestamp | Confidence

THE EXCEL IS A COPY OF THE TERMINAL:
  Every "[step] ..." line printed on the terminal is appended to the Excel, in the
  same order. Nothing printed as a [step] line is ever removed or rewritten later.
  The periodic progress lines ("[123/4567] ...") are NOT logged.

Labels are taken directly from the model's config.json
(step1 ... step9, sop_step12). No renaming is applied unless you edit LABEL_RENAME.

Each time a step ends (the predicted label changes) the new steps are printed and
appended to the Excel. Press Ctrl+C at any time: the cache and the Excel are saved
(including the step currently in progress) and the program exits.

RESUME CACHE
  Every window's raw prediction (class id + confidence) AND the step log are stored
  in a cache file (default: <xlsx_out>.cache.json). If you stop, pause (Ctrl+C),
  crash, or the PC shuts down, run the SAME command again: the program loads the
  cache and continues from the first window that has not been processed yet.
  - The cache is written every --cache_every windows, on every logged step, and on
    Ctrl+C / SIGTERM / error.
  - The cache is only reused if the video, model, ROI, window_len, stride and
    num_frames are unchanged. Otherwise it is ignored (and a warning is shown).
  - If post-processing options change between runs, the step log is rebuilt from the
    cached RAW predictions (replay), so you never need to redo inference.
  - Use --fresh to ignore the cache and start from the beginning.

Post-processing (used to decide which steps get printed/logged; every part can be switched off):
  1. window argmax with confidence < --confidence_threshold  -> "uncertain"
  2. consecutive same-label windows merged, short segments absorbed (--min_segment_duration)
  3. "uncertain" segments dropped
  4. segment confidence = mean confidence of ALL windows whose midpoint lies inside it
  5. overlapping same-label segments merged (end = max, conf = max)
  6. segments with confidence < --seg_conf dropped
  7. step-order filter: keeps the best order-non-decreasing chain (max sum conf*duration)

Usage:
  python log_videomae_detections.py \
      --videomae_dir runs/roi_v8/videomae_finetuned/best_model \
      --long_video caa_16hrs_trimmed.mkv \
      --xlsx_out detections.xlsx \
      --window_len 6.5 --stride 1.0 --fps 15 --num_frames 16

Optional:
  --cache path/to/cache.json   # default: <xlsx_out>.cache.json
  --cache_every 25             # save cache every N windows
  --fresh                      # ignore existing cache, start over
  --roi 340,201,413,439   # x,y,width,height; otherwise reads model_dir/roi.json
  --confidence_threshold 0.7
  --seg_conf 0.7
  --min_segment_duration 2.5
  --no_postproc / --no_merge / --no_conf_filter / --no_order_filter
  --order step1,step2,...,sop_step12
  --cpu
  --fp16

Dependencies:
  pip install torch transformers opencv-python numpy openpyxl
"""

import argparse
import json
import os
import signal
import sys
from pathlib import Path

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

UNCERTAIN = "uncertain"
CACHE_VERSION = 2

# Optional renaming of model labels in the Excel, e.g. {"sop_step12": "stop_step12"}.
# Default: empty = names exactly as in the model's config.json.
LABEL_RENAME = {}

# Order used by the step-order filter (names as the model outputs them).
DEFAULT_ORDER = ["step1", "step2", "step3", "step4", "step5", "step6",
                 "step7", "step8", "step9", "sop_step12"]


def hms(seconds: float) -> str:
    seconds = max(0.0, float(seconds))
    ms = int(round(seconds * 1000))
    hours, ms = divmod(ms, 3_600_000)
    minutes, ms = divmod(ms, 60_000)
    secs, ms = divmod(ms, 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}.{ms:03d}"


class VideoMAE:
    def __init__(self, model_dir, use_cuda=True, fp16=False):
        import torch
        from transformers import VideoMAEForVideoClassification

        self.torch = torch
        self.resize_short, self.crop_h, self.crop_w = 224, 224, 224
        self.mean = np.array([0.485, 0.456, 0.406], np.float32)
        self.std = np.array([0.229, 0.224, 0.225], np.float32)

        preproc = Path(model_dir) / "preprocessor_config.json"
        if preproc.exists():
            cfg = json.loads(preproc.read_text())

            def size_value(value, default, first, second):
                if isinstance(value, (int, float)):
                    return int(value)
                if isinstance(value, dict):
                    if first in value:
                        return int(value[first])
                    if second in value:
                        return int(value[second])
                return default

            if "size" in cfg:
                self.resize_short = size_value(cfg["size"], 224, "shortest_edge", "height")
            if "crop_size" in cfg:
                self.crop_h = size_value(cfg["crop_size"], 224, "height", "shortest_edge")
                self.crop_w = size_value(cfg["crop_size"], 224, "width", "shortest_edge")
            if len(cfg.get("image_mean", [])) == 3:
                self.mean = np.array(cfg["image_mean"], np.float32)
            if len(cfg.get("image_std", [])) == 3:
                self.std = np.array(cfg["image_std"], np.float32)

        if use_cuda and not torch.cuda.is_available():
            print("[warn] CUDA unavailable in this PyTorch; using CPU.", file=sys.stderr)
        self.device = "cuda" if use_cuda and torch.cuda.is_available() else "cpu"
        self.fp16 = bool(fp16 and self.device == "cuda")
        self.model = VideoMAEForVideoClassification.from_pretrained(model_dir).eval().to(self.device)
        if self.fp16:
            self.model.half()
        print(f"[model] device={self.device}, fp16={self.fp16}")

    def preprocess(self, frames):
        height, width = self.crop_h, self.crop_w
        batch = np.empty((len(frames), 3, height, width), np.float32)
        for i, rgb in enumerate(frames):
            h, w = rgb.shape[:2]
            short, long = min(h, w), max(h, w)
            new_short = self.resize_short
            new_long = int(new_short * long / short)
            new_w, new_h = (new_short, new_long) if w <= h else (new_long, new_short)
            interp = cv2.INTER_AREA if (new_w < w or new_h < h) else cv2.INTER_LINEAR
            resized = cv2.resize(rgb, (new_w, new_h), interpolation=interp)

            if resized.shape[0] < height or resized.shape[1] < width:
                padded = np.zeros(
                    (max(height, resized.shape[0]), max(width, resized.shape[1]), 3),
                    dtype=np.uint8,
                )
                padded[:resized.shape[0], :resized.shape[1]] = resized
                resized = padded

            top = (resized.shape[0] - height) // 2
            left = (resized.shape[1] - width) // 2
            crop = resized[top:top + height, left:left + width].astype(np.float32) / 255.0
            batch[i] = ((crop - self.mean) / self.std).transpose(2, 0, 1)
        return batch[None]

    def predict(self, frames):
        torch = self.torch
        try:
            x = torch.from_numpy(self.preprocess(frames)).to(self.device)
            if self.fp16:
                x = x.half()
            with torch.inference_mode():
                logits = self.model(pixel_values=x).logits
            logits = logits[0].float().cpu().numpy().astype(np.float64)
        except torch.cuda.OutOfMemoryError as exc:
            torch.cuda.empty_cache()
            raise RuntimeError("GPU out of memory. Try --fp16 or --cpu.") from exc

        logits -= logits.max()
        probs = np.exp(logits)
        return probs / probs.sum()


def read_frames(cap, start_sec, end_sec, native_fps, total_frames, num_frames, roi, video_path):
    start_frame = int(start_sec * native_fps)
    end_frame = max(int(end_sec * native_fps), start_frame + 1)
    indexes = []
    for i in range(num_frames):
        frame_idx = (
            start_frame if num_frames == 1
            else start_frame + (end_frame - 1 - start_frame) * i / (num_frames - 1)
        )
        indexes.append(int(min(max(round(frame_idx), 0), max(total_frames - 1, 0))))

    frames = [None] * num_frames
    cap.set(cv2.CAP_PROP_POS_FRAMES, indexes[0])
    k, last = 0, None
    for frame_idx in range(indexes[0], indexes[-1] + 1):
        if k >= len(indexes) or not cap.grab():
            break
        if indexes[k] != frame_idx:
            continue
        ok, bgr = cap.retrieve()
        if not ok or bgr is None:
            break

        rgb = cv2.cvtColor(bgr, cv2.COLOR_BGR2RGB)
        if roi:
            x, y, w, h = roi
            x0, y0 = max(0, min(x, rgb.shape[1])), max(0, min(y, rgb.shape[0]))
            x1 = max(0, min(x + w, rgb.shape[1]))
            y1 = max(0, min(y + h, rgb.shape[0]))
            if x1 > x0 and y1 > y0:
                rgb = rgb[y0:y1, x0:x1]

        last = np.ascontiguousarray(rgb)
        while k < len(indexes) and indexes[k] == frame_idx:
            frames[k] = last
            k += 1

    if last is None:
        raise RuntimeError(f"Could not decode frames from {video_path}")
    return [frame if frame is not None else last for frame in frames]


def resolve_roi(model_dir, roi_arg):
    if roi_arg:
        try:
            roi = tuple(int(v.strip()) for v in roi_arg.split(","))
        except ValueError as exc:
            raise SystemExit("--roi must be x,y,width,height") from exc
        if len(roi) != 4:
            raise SystemExit("--roi must be x,y,width,height")
        return roi

    roi_path = Path(model_dir) / "roi.json"
    if roi_path.exists():
        data = json.loads(roi_path.read_text())
        roi = data.get("roi")
        if roi and len(roi) == 4 and all(isinstance(v, (int, float)) for v in roi):
            result = tuple(int(v) for v in roi)
            print(f"[roi] Auto-loaded {result} from {roi_path}")
            return result
    return None


# ----------------------------------------------------------------------------- cache
def model_fingerprint(model_dir):
    """Identify the model weights (name, size) so a retrained model invalidates the cache."""
    model_dir = Path(model_dir)
    for name in ("model.safetensors", "pytorch_model.bin"):
        p = model_dir / name
        if p.exists():
            st = p.stat()
            return {"file": name, "size": st.st_size, "mtime": int(st.st_mtime)}
    return {"file": None}


def build_signature(args, video_path, roi, native_fps, total_frames):
    """Everything that changes the RAW per-window predictions."""
    return {
        "video_name": os.path.basename(video_path),
        "video_size": os.path.getsize(video_path),
        "total_frames": int(total_frames),
        "native_fps": round(float(native_fps), 4),
        "window_len": float(args.window_len),
        "stride": float(args.stride),
        "num_frames": int(args.num_frames),
        "roi": list(roi) if roi else None,
        "model": model_fingerprint(args.videomae_dir),
    }


def build_post_signature(args):
    """Everything that changes WHICH steps get logged (post-processing settings)."""
    return {
        "confidence_threshold": float(args.confidence_threshold),
        "seg_conf": float(args.seg_conf),
        "min_segment_duration": float(args.min_segment_duration),
        "postproc": bool(args.postproc),
        "merge": bool(args.merge),
        "conf_filter": bool(args.conf_filter),
        "order_filter": bool(args.order_filter),
        "order": list(args.order),
    }


def load_cache(path, signature, num_labels, total_windows):
    """Return (class_ids, confidences, log, post_sig) from a valid cache, or ([], [], None, None)."""
    empty = ([], [], None, None)
    path = Path(path)
    if not path.exists():
        return empty
    try:
        data = json.loads(path.read_text())
    except (OSError, ValueError) as exc:
        print(f"[cache] Could not read {path} ({exc}); starting from the beginning.")
        return empty

    if data.get("version") != CACHE_VERSION:
        print("[cache] Cache version mismatch; starting from the beginning.")
        return empty
    if data.get("signature") != signature:
        diff = [k for k in signature if data.get("signature", {}).get(k) != signature[k]]
        print(f"[cache] Cache does not match the current run (changed: {', '.join(diff)}); "
              f"ignoring it and starting from the beginning.")
        return empty

    class_ids = data.get("class_ids", [])
    confidences = data.get("confidences", [])
    if (len(class_ids) != len(confidences) or len(class_ids) > total_windows
            or any((not isinstance(c, int)) or c < 0 or c >= num_labels for c in class_ids)):
        print("[cache] Cache content is inconsistent; starting from the beginning.")
        return empty

    log = None
    raw_log = data.get("log")
    if isinstance(raw_log, list):
        try:
            log = [(str(e[0]), float(e[1]), float(e[2]), float(e[3])) for e in raw_log]
        except (IndexError, TypeError, ValueError):
            log = None
    return class_ids, confidences, log, data.get("post_sig")


def save_cache(path, signature, post_sig, class_ids, confidences, log):
    """Atomic write so a crash/power loss during saving never corrupts the cache."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + ".tmp")
    payload = {
        "version": CACHE_VERSION,
        "signature": signature,
        "post_sig": post_sig,
        "class_ids": [int(c) for c in class_ids],
        "confidences": [round(float(c), 6) for c in confidences],
        "log": [[l, float(s), float(e), float(c)] for l, s, e, c in log],
    }
    try:
        temp.write_text(json.dumps(payload))
        os.replace(temp, path)
    except OSError as exc:
        print(f"[warn] Could not save cache {path}: {exc}", file=sys.stderr)


# ----------------------------------------------------------------------------- segmenting
def merge_predictions(windows, labels, confidences, min_duration):
    """Merge consecutive same-label windows into temporal segments (label, start, end, mean_conf)."""
    if not labels:
        return []

    segments = []
    current_label = labels[0]
    start = windows[0][0]
    end = windows[0][1]
    conf_values = [confidences[0]]

    for i in range(1, len(labels)):
        if labels[i] == current_label:
            end = max(end, windows[i][1])
            conf_values.append(confidences[i])
        else:
            segments.append((current_label, start, end, float(np.mean(conf_values))))
            current_label = labels[i]
            start, end = windows[i]
            conf_values = [confidences[i]]
    segments.append((current_label, start, end, float(np.mean(conf_values))))

    # Merge short segments into the longer adjacent segment. 0 disables this.
    if min_duration > 0:
        changed = True
        while changed and len(segments) > 1:
            changed = False
            for i, seg in enumerate(segments):
                if seg[2] - seg[1] >= min_duration:
                    continue
                left_duration = segments[i - 1][2] - segments[i - 1][1] if i > 0 else -1
                right_duration = segments[i + 1][2] - segments[i + 1][1] if i + 1 < len(segments) else -1
                if left_duration == -1 and right_duration == -1:
                    continue
                if left_duration >= right_duration:
                    prev = segments[i - 1]
                    segments[i - 1] = (prev[0], prev[1], seg[2], prev[3])
                else:
                    nxt = segments[i + 1]
                    segments[i + 1] = (nxt[0], seg[1], nxt[2], nxt[3])
                del segments[i]
                changed = True
                break
    return segments


def segment_confidence(windows, confidences, start, end):
    """Mean window confidence over ALL windows whose midpoint lies inside [start, end]."""
    if not windows:
        return 0.0
    mids = np.array([(w[0] + w[1]) / 2 for w in windows])
    conf = np.asarray(confidences, dtype=np.float64)
    m = (mids >= start) & (mids <= end)
    return float(conf[m].mean()) if m.any() else float(conf.mean())


def order_filter(dets, order):
    """Keep the best order-non-decreasing chain (max sum of conf*duration).
    Ties: later predecessor wins, later chain end wins. Labels not in `order` are dropped."""
    oidx = lambda l: order.index(l) if l in order else -1
    cand = sorted((d for d in dets if oidx(d["label"]) >= 0), key=lambda d: (d["start"], d["end"]))
    n = len(cand)
    if n == 0:
        return []
    w = [d["conf"] * (d["end"] - d["start"]) for d in cand]
    oi = [oidx(d["label"]) for d in cand]
    best, prev, EPS = [0.0] * n, [-1] * n, 1e-12
    for i in range(n):
        bj, bv = -1, 0.0
        for j in range(i):
            if oi[j] > oi[i]:
                continue
            if best[j] > bv + EPS or (bj >= 0 and abs(best[j] - bv) <= EPS and best[j] > 0):
                bv, bj = best[j], j
        best[i] = w[i] + bv
        prev[i] = bj
    e = 0
    for i in range(1, n):
        if best[i] >= best[e] - EPS:
            e = i
    chain, i = [], e
    while i >= 0:
        chain.append(cand[i])
        i = prev[i]
    chain.reverse()
    return chain


def build_detections(windows, labels, confidences, args):
    """Full pipeline: windows -> list of (label, start, end, confidence)."""
    if not labels:
        return []
    segs = merge_predictions(windows, labels, confidences, args.min_segment_duration)
    dets = [{"label": s[0], "start": s[1], "end": s[2],
             "conf": segment_confidence(windows, confidences, s[1], s[2])}
            for s in segs if s[0] != UNCERTAIN]

    if args.postproc:
        if args.merge and dets:
            dets.sort(key=lambda d: d["start"])
            merged = []
            for d in dets:
                if merged and merged[-1]["label"] == d["label"] and d["start"] <= merged[-1]["end"]:
                    m = merged[-1]
                    m["end"] = max(m["end"], d["end"])
                    m["conf"] = max(m["conf"], d["conf"])
                else:
                    merged.append(dict(d))
            dets = merged
        if args.conf_filter:
            dets = [d for d in dets if d["conf"] >= args.seg_conf]
        if args.order_filter and dets:
            dets = order_filter(dets, args.order)

    return [(d["label"], d["start"], d["end"], d["conf"]) for d in dets]


# ----------------------------------------------------------------------------- step log
def step_key(label, start, end):
    return (label, round(start, 2), round(end, 2))


def log_new_steps(detections, printed, log, quiet=False):
    """Print every step not yet shown AND append it to the log (= future Excel rows).
    The log is append-only: whatever was printed stays in the Excel forever.
    Returns the number of newly logged steps."""
    added = 0
    for label, start, end, conf in detections:
        key = step_key(label, start, end)
        if key in printed:
            continue
        printed.add(key)
        log.append((label, float(start), float(end), float(conf)))
        added += 1
        if not quiet:
            print(f"[step] {label:<12} {hms(start)} - {hms(end)}  conf={conf:.4f}")
    return added


def replay_log(windows, labels, confidences, args, quiet=True):
    """Rebuild the step log from raw predictions exactly as the live loop would have made it."""
    log, printed = [], set()
    n = len(labels)
    for i in range(2, n + 1):
        if labels[i - 1] != labels[i - 2]:
            dets = build_detections(windows[:i - 1], labels[:i - 1], confidences[:i - 1], args)
            log_new_steps(dets, printed, log, quiet=quiet)
    return log, printed


# ----------------------------------------------------------------------------- excel
def write_excel(path, log):
    """Write the step log (exactly the printed [step] lines, in order) to Excel."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Detections"
    headers = ["Detection", "Video Timestamp", "Confidence"]
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for label, start, end, confidence in log:
        sheet.append([label, f"{hms(start)} - {hms(end)}", round(confidence, 4)])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col, width in {"A": 28, "B": 38, "C": 16}.items():
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(horizontal="left")
        row[2].number_format = "0.0000"

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = output.with_name(output.stem + ".tmp.xlsx")
    workbook.save(temp)
    try:
        os.replace(temp, output)
    except PermissionError:
        alternate = output.with_name(f"{output.stem}_new.xlsx")
        os.replace(temp, alternate)
        print(f"[warn] Output file was locked (open in Excel?); saved as {alternate}")
        return alternate
    return output


# ----------------------------------------------------------------------------- args / main
def parse_args():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--videomae_dir", required=True, help="Path to VideoMAE best_model directory")
    parser.add_argument("--long_video", required=True, help="Input video path")
    parser.add_argument("--xlsx_out", default="detections.xlsx", help="Output Excel path")
    parser.add_argument("--cache", default="",
                        help="Resume-cache file (default: <xlsx_out>.cache.json)")
    parser.add_argument("--cache_every", type=int, default=25,
                        help="Save the resume cache every N processed windows (default 25)")
    parser.add_argument("--fresh", action="store_true",
                        help="Ignore any existing cache and start from the beginning")
    parser.add_argument("--window_len", type=float, default=6.5)
    parser.add_argument("--stride", type=float, default=1.0)
    parser.add_argument("--fps", type=float, default=15.0, help="Fallback FPS if video metadata has no FPS")
    parser.add_argument("--num_frames", type=int, default=16)
    parser.add_argument("--min_segment_duration", type=float, default=2.5)
    parser.add_argument("--confidence_threshold", type=float, default=0.7,
                        help="Windows whose top-class probability is below this become 'uncertain' "
                             "(default 0.7, as in make_test_report.py; 0 = no threshold)")
    parser.add_argument("--seg_conf", type=float, default=0.7,
                        help="Drop segments with confidence below this (default 0.7)")
    parser.add_argument("--no_postproc", action="store_true", help="Turn all post-processing off (raw segments)")
    parser.add_argument("--no_merge", action="store_true", help="Skip merging of overlapping same-label segments")
    parser.add_argument("--no_conf_filter", action="store_true", help="Skip the segment-confidence filter")
    parser.add_argument("--no_order_filter", action="store_true", help="Skip the step-order filter")
    parser.add_argument("--order", default=",".join(DEFAULT_ORDER),
                        help="Comma separated step order used by the order filter")
    parser.add_argument("--roi", default="", help="ROI as x,y,width,height; otherwise read roi.json from model directory")
    parser.add_argument("--cpu", action="store_true", help="Run inference on CPU")
    parser.add_argument("--fp16", action="store_true", help="Use half precision on CUDA")
    a = parser.parse_args()
    a.postproc = not a.no_postproc
    a.merge = not a.no_merge
    a.conf_filter = not a.no_conf_filter
    a.order_filter = not a.no_order_filter
    a.order = [x.strip() for x in a.order.split(",") if x.strip()]
    if not a.cache:
        a.cache = str(Path(a.xlsx_out).with_name(Path(a.xlsx_out).name + ".cache.json"))
    a.cache_every = max(1, a.cache_every)
    return a


def main():
    args = parse_args()
    if args.window_len <= 0 or args.stride <= 0 or args.num_frames < 1:
        raise SystemExit("--window_len and --stride must be > 0, and --num_frames must be >= 1")
    if not 0 <= args.confidence_threshold <= 1:
        raise SystemExit("--confidence_threshold must be between 0 and 1")

    model_dir = Path(args.videomae_dir)
    config_path = model_dir / "config.json"
    if not config_path.exists():
        raise SystemExit(f"Cannot find model config: {config_path}")
    model_config = json.loads(config_path.read_text())
    id2label = model_config.get("id2label", {})
    labels_by_id = [str(id2label.get(str(i), id2label.get(i, f"class_{i}"))).strip()
                    for i in range(len(id2label))]
    if not labels_by_id:
        raise SystemExit(f"No id2label entries found in {config_path}")
    labels_by_id = [LABEL_RENAME.get(l, l) for l in labels_by_id]

    # Make sure the order filter never silently drops a label it doesn't know about.
    unknown = [l for l in dict.fromkeys(labels_by_id) if l not in args.order]
    if unknown and args.postproc and args.order_filter:
        print(f"[warn] model labels not in the step order list: {unknown} -> appended at the end of the order. "
              f"Use --order to set the correct position.")
        args.order = args.order + unknown

    roi = resolve_roi(model_dir, args.roi)
    video_path = args.long_video
    if not os.path.isfile(video_path):
        raise SystemExit(f"Cannot find video: {video_path}")

    cap = cv2.VideoCapture(video_path)
    if not cap.isOpened():
        raise SystemExit(f"Cannot open video: {video_path}")

    total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
    native_fps = cap.get(cv2.CAP_PROP_FPS)
    if not native_fps or native_fps <= 0:
        native_fps = args.fps
    duration = total_frames / native_fps if total_frames > 0 else 0
    if duration <= 0:
        cap.release()
        raise SystemExit("Video duration could not be determined.")

    windows = []
    t = 0.0
    while t + args.window_len <= duration + 1e-6:
        windows.append((t, min(t + args.window_len, duration)))
        t += args.stride
    if not windows:
        windows.append((0.0, duration))

    print(f"[video] {video_path}")
    print(f"[video] duration={hms(duration)}, native_fps={native_fps:.3f}, windows={len(windows)}")
    print(f"[labels] {labels_by_id}")
    if roi:
        print(f"[roi] {roi}")
    print("[info] Press Ctrl+C at any time: the cache + Excel are saved and the program exits. "
          "Run the same command again to resume.")

    # ------------------------------------------------------------------ resume from cache
    signature = build_signature(args, video_path, roi, native_fps, total_frames)
    post_sig = build_post_signature(args)
    class_ids, confidences, cached_log, cached_post = [], [], None, None
    if args.fresh:
        print("[cache] --fresh given: ignoring any existing cache.")
    else:
        class_ids, confidences, cached_log, cached_post = load_cache(
            args.cache, signature, len(labels_by_id), len(windows))

    def to_label(class_id, confidence):
        return UNCERTAIN if confidence < args.confidence_threshold else labels_by_id[class_id]

    # Labels are recomputed from raw predictions, so --confidence_threshold may change between runs.
    predicted_labels = [to_label(c, p) for c, p in zip(class_ids, confidences)]
    n_done = len(class_ids)

    log = []          # append-only list of every [step] line printed (= Excel rows)
    printed = set()   # keys of steps already printed/logged

    if n_done:
        if n_done >= len(windows):
            print(f"[cache] Loaded {n_done}/{len(windows)} windows: video was already fully processed "
                  f"(use --fresh to redo it).")
        else:
            print(f"[cache] Resuming: {n_done}/{len(windows)} windows already done "
                  f"(up to {hms(windows[n_done - 1][1])}); continuing from {hms(windows[n_done][0])}.")

        if cached_log is not None and cached_post == post_sig:
            log = list(cached_log)
            printed = {step_key(l, s, e) for l, s, e, _ in log}
        else:
            print("[cache] Post-processing settings changed (or no step log in cache): "
                  "rebuilding the step log from cached predictions...")
            log, printed = replay_log(windows, predicted_labels, confidences, args)

        for label, start, end, conf in log:
            print(f"[step] {label:<12} {hms(start)} - {hms(end)}  conf={conf:.4f}")
    else:
        print(f"[cache] Starting from the beginning (cache file: {args.cache})")

    # Make SIGTERM (kill / task manager stop on Linux/macOS) behave like Ctrl+C so state is saved.
    if hasattr(signal, "SIGTERM"):
        def _on_sigterm(signum, frame):
            raise KeyboardInterrupt
        try:
            signal.signal(signal.SIGTERM, _on_sigterm)
        except (ValueError, OSError):
            pass

    interrupted = False
    output = write_excel(args.xlsx_out, log)

    model = None
    if n_done < len(windows):
        model = VideoMAE(args.videomae_dir, use_cuda=not args.cpu, fp16=args.fp16)

    try:
        for i in range(n_done + 1, len(windows) + 1):
            start, end = windows[i - 1]
            frames = read_frames(cap, start, end, native_fps, total_frames,
                                 args.num_frames, roi, video_path)
            probabilities = model.predict(frames)
            if len(probabilities) != len(labels_by_id):
                raise RuntimeError(
                    f"Model output has {len(probabilities)} classes but "
                    f"config.json has {len(labels_by_id)} labels")

            class_id = int(np.argmax(probabilities))
            confidence = float(probabilities[class_id])
            label = to_label(class_id, confidence)

            # Window fully processed -> record it (lists always stay the same length).
            class_ids.append(class_id)
            confidences.append(confidence)
            predicted_labels.append(label)

            # A step has just ENDED when the label changes -> print + log everything new.
            # (windows[:i-1] excludes the window that started the new label.)
            if i > 1 and label != predicted_labels[-2]:
                detections = build_detections(windows[:i - 1], predicted_labels[:i - 1],
                                              confidences[:i - 1], args)
                if log_new_steps(detections, printed, log):
                    output = write_excel(args.xlsx_out, log)
                    save_cache(args.cache, signature, post_sig, class_ids, confidences, log)
                elif i % args.cache_every == 0:
                    save_cache(args.cache, signature, post_sig, class_ids, confidences, log)
            elif i % args.cache_every == 0:
                save_cache(args.cache, signature, post_sig, class_ids, confidences, log)

            # Progress line: console only, NEVER logged to Excel.
            if i == 1 or i % 100 == 0 or i == len(windows):
                print(f"[{i}/{len(windows)}] {hms(start)} -> {label} ({confidence:.4f})")

    except KeyboardInterrupt:
        interrupted = True
        # Ignore further Ctrl+C presses so the final save cannot be cut off.
        signal.signal(signal.SIGINT, signal.SIG_IGN)
        print("\n[interrupt] Stop received. Saving cache and Excel...")

    finally:
        try:
            signal.signal(signal.SIGINT, signal.SIG_IGN)
        except Exception:
            pass
        cap.release()
        # Final flush: print + log the steps still pending (including the step currently
        # in progress), then write the Excel, then save the cache (with the same log).
        if predicted_labels:
            final_dets = build_detections(windows[:len(predicted_labels)], predicted_labels,
                                          confidences, args)
            log_new_steps(final_dets, printed, log)
        output = write_excel(args.xlsx_out, log)
        save_cache(args.cache, signature, post_sig, class_ids, confidences, log)

    done = len(predicted_labels)
    if interrupted:
        print(f"Stopped. Steps logged: {len(log)}")
        print(f"Excel: {output}")
        print(f"Cache: {args.cache}")
        print(f"Processed {done}/{len(windows)} windows, up to "
              f"{hms(windows[done - 1][1]) if done else hms(0)} of the video. "
              "Run the same command again to RESUME from this point (add --fresh to restart).")
        sys.exit(0)
    else:
        print(f"\nDone. Steps logged: {len(log)}")
        print(f"Excel: {output}")
        print(f"Cache: {args.cache}  (re-running resumes/finishes instantly; use --fresh to redo)")


if __name__ == "__main__":
    main()
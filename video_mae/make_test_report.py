#!/usr/bin/env python3
"""
Runs the VideoMAE temporal test and fills temporal_test_template.xlsx.

  Step 1 (individual clips)  : every video in --individual_dir; true label is parsed
                               from the filename (start_step1_op1_16.mkv -> start_step1).
  Step 2 (continuous videos) : every cycleN.mkv in --continuous_dir; ground truth is read
                               from cycle20GT.ods (row for cycle N). One Excel sheet per cycle.

It uses the SAME functions/settings as detect_segments.py (extract_session_embeddings,
labels_to_segments, apply_min_duration_filter, confidence threshold), but calls them directly
so it can also record per-clip confidence.

USAGE (from the video_mae folder, same env you use for detect_segments.py):

    python3 make_test_report.py \
        --videomae_dir runs/roi_v4/videomae_finetuned/best_model \
        --individual_dir videos/ \
        --continuous_dir caa_20cycles/ \
        --gt cycle20GT.ods \
        --template temporal_test_template.xlsx \
        --window_len 2.0 --stride 1.0 --fps 15 --num_frames 16 \
        --min_segment_duration 2.5 --confidence_threshold 0.7
"""

import argparse
import json
import os
import re
import sys
from copy import copy
from datetime import datetime

import numpy as np

UNCERTAIN = "uncertain"
MERGE_STEP6_RENAME = {"step6": "step6.5"}   # only used with --merge_step6 (default: step6 and step6.5 are separate classes)
GT_MERGE_GAP = 3.0                      # GT ranges with the same label closer than this get merged
MAX_ROWS = 40                           # template capacity for GT rows and detected rows per sheet


# ----------------------------------------------------------------------------- labels
_STEP_RE = re.compile(r"^(?:(start|stop)_)?step_?(\d+(?:\.\d+)?)$", re.I)
_FILE_RE = re.compile(r"^(?P<lab>(?:(?:start|stop)_)?step_?\d+(?:\.\d+)?)_(?:op\d+_)?\d+$", re.I)


def canon(name, rename):
    """Normalise any spelling of a step name to the Excel naming."""
    s = str(name).strip().lower().replace(" ", "")
    m = _STEP_RE.match(s)
    if m:
        pre, n = m.groups()
        if pre == "stop":
            s = f"stop_step{n}"
        elif n in ("1", "2"):
            s = f"start_step{n}"
        else:
            s = f"step{n}"
    return rename.get(s, s)


def label_from_filename(path, rename):
    stem = os.path.splitext(os.path.basename(path))[0]
    m = _FILE_RE.match(stem)
    return canon(m.group("lab"), rename) if m else None


# ----------------------------------------------------------------------------- ground truth
def _to_sec(v):
    import datetime as dt
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    if isinstance(v, dt.time):
        return v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, dt.timedelta):
        return v.total_seconds()
    if isinstance(v, (int, float)):
        return float(v)
    parts = [float(p) for p in str(v).strip().split(":")]
    sec = 0.0
    for p in parts:
        sec = sec * 60 + p
    return sec


def load_gt(path, rename):
    """-> {cycle_no: [ {label,start,end}, ... ] sorted by start}"""
    import pandas as pd
    engine = "odf" if path.lower().endswith(".ods") else None
    df = pd.read_excel(path, engine=engine, header=None)
    hdr, sub = df.iloc[0], df.iloc[1]
    cyc_col = next(c for c in range(df.shape[1]) if str(hdr[c]).strip().lower() == "cycle")

    step_cols = []
    for c in range(cyc_col + 1, df.shape[1] - 1):
        h = hdr[c]
        if isinstance(h, str) and h.strip() and str(sub[c]).strip().lower().startswith("from") \
                and str(sub[c + 1]).strip().lower() == "to":
            name = re.split(r"\s*[\u2013\u2014-]\s*", h.strip(), maxsplit=1)[0]
            step_cols.append((canon(name, rename), c))

    gt = {}
    for r in range(2, df.shape[0]):
        cyc = df.iloc[r, cyc_col]
        if cyc is None or (isinstance(cyc, float) and np.isnan(cyc)):
            continue
        cyc = int(cyc)
        segs = []
        for lab, c in step_cols:
            s, e = _to_sec(df.iloc[r, c]), _to_sec(df.iloc[r, c + 1])
            if s is None or e is None or e <= s:
                print(f"  [GT warning] cycle {cyc} {lab}: bad/missing times ({df.iloc[r, c]} -> {df.iloc[r, c + 1]}), skipped")
                continue
            segs.append({"label": lab, "start": s, "end": e})
        # merge same-label ranges (happens when step6 + step6.5 are merged into one label)
        segs.sort(key=lambda x: (x["label"], x["start"]))
        merged = []
        for sg in segs:
            if merged and merged[-1]["label"] == sg["label"] and sg["start"] <= merged[-1]["end"] + GT_MERGE_GAP:
                merged[-1]["end"] = max(merged[-1]["end"], sg["end"])
            else:
                merged.append(dict(sg))
        merged.sort(key=lambda x: x["start"])
        gt[cyc] = merged
    return gt, [lab for lab, _ in step_cols]


# ----------------------------------------------------------------------------- inference
class Engine:
    """Lazy model loader + per-video (windows, probs) with an on-disk cache."""

    def __init__(self, args, be, class_names_raw, roi, fingerprint):
        self.a, self.be, self.roi = args, be, roi
        self.model = self.processor = None
        key = f"{fingerprint}_w{args.window_len}_s{args.stride}_f{args.fps}_n{args.num_frames}_roi{'-'.join(map(str, roi)) if roi else 'none'}"
        self.cache_dir = None if args.no_cache else os.path.join(args.runs_dir, "report_cache", key)
        if self.cache_dir:
            os.makedirs(self.cache_dir, exist_ok=True)

    def _load(self):
        if self.model is None:
            from transformers import VideoMAEForVideoClassification, VideoMAEImageProcessor
            self.processor = VideoMAEImageProcessor.from_pretrained(self.a.videomae_dir)
            self.model = VideoMAEForVideoClassification.from_pretrained(self.a.videomae_dir)

    def windows_and_probs(self, video_path):
        cpath = None
        if self.cache_dir:
            st = os.stat(video_path)
            cpath = os.path.join(self.cache_dir, f"{os.path.splitext(os.path.basename(video_path))[0]}_{st.st_size}.npz")
            if os.path.exists(cpath):
                z = np.load(cpath)
                return [tuple(w) for w in z["windows"]], z["probs"]
        self._load()
        a = self.a
        window_times, _emb, probs, timing = self.be.extract_session_embeddings(
            video_path, self.model, self.processor,
            window_len=a.window_len, stride=a.stride, fps=a.fps, num_frames=a.num_frames, roi=self.roi)
        probs = np.asarray(probs, dtype=np.float64)
        if cpath:
            np.savez(cpath, windows=np.array(window_times, dtype=np.float64), probs=probs)
        return window_times, probs


def collapse_matrix(class_names_raw, rename):
    names = [canon(n, rename) for n in class_names_raw]
    uniq = list(dict.fromkeys(names))
    M = np.zeros((len(names), len(uniq)))
    for i, n in enumerate(names):
        M[i, uniq.index(n)] = 1.0
    return uniq, M


def windows_to_segments(window_times, probs_c, uniq, a, be):
    if len(window_times) == 0:
        return [], [], np.array([])
    starts = [w[0] for w in window_times]
    ends = [w[1] for w in window_times]
    idx, conf = probs_c.argmax(axis=1), probs_c.max(axis=1)
    labels = [uniq[i] for i in idx]
    if a.confidence_threshold is not None:
        labels = [l if c >= a.confidence_threshold else UNCERTAIN for l, c in zip(labels, conf)]
    raw = be.labels_to_segments(starts, ends, labels)
    segs = be.apply_min_duration_filter(raw, a.min_segment_duration)
    return segs, labels, conf


def clip_prediction(window_times, probs_c, uniq, segs, conf, a):
    """One (label, confidence, note) for a single-class clip."""
    if len(window_times) == 0:
        return UNCERTAIN, 0.0, "no windows (clip shorter than window_len?)"
    if a.clip_method == "mean_prob":
        mp = probs_c.mean(axis=0)
        i = int(mp.argmax())
        lab, c = uniq[i], float(mp[i])
        if a.confidence_threshold is not None and c < a.confidence_threshold:
            lab = UNCERTAIN
        note = ""
    else:  # dominant_segment: real label with the most time; 'uncertain' only if NO real label was detected
        dur = {}
        for s in segs:
            dur[s["label"]] = dur.get(s["label"], 0.0) + (s["end_sec"] - s["start_sec"])
        real = {k: v for k, v in dur.items() if k != UNCERTAIN}
        lab = max(real, key=real.get) if real else UNCERTAIN
        mids = np.array([(w[0] + w[1]) / 2 for w in window_times])
        mask = np.zeros(len(mids), bool)
        for s in segs:
            if s["label"] == lab:
                mask |= (mids >= s["start_sec"]) & (mids <= s["end_sec"])
        c = float(conf[mask].mean()) if mask.any() else float(conf.mean())
        note = ""
    if len(segs) > 1:
        note = "segments: " + "; ".join(f"{s['label']} {s['start_sec']:.0f}-{s['end_sec']:.0f}s" for s in segs)
    return lab, c, note


# ----------------------------------------------------------------------------- matching (Section C)
def match_gt_to_det(gt_segs, det_segs, min_overlap):
    """For each GT segment return 1-based index of a detected segment (or None = missed).
    Prefers a same-label detection with the biggest overlap; else biggest-overlap other label."""
    def ov(g, d):
        return max(0.0, min(g["end"], d["end"]) - max(g["start"], d["start"]))
    out = []
    for g in gt_segs:
        same = [(ov(g, d), i) for i, d in enumerate(det_segs) if d["label"] == g["label"] and ov(g, d) >= min_overlap]
        other = [(ov(g, d), i) for i, d in enumerate(det_segs) if d["label"] != g["label"] and ov(g, d) >= min_overlap]
        pick = max(same)[1] if same else (max(other)[1] if other else None)
        out.append(None if pick is None else pick + 1)
    return out


# ----------------------------------------------------------------------------- excel helpers
def _reset(cell):
    from openpyxl.styles.cell_style import StyleArray
    cell.value = None
    cell._style = StyleArray()


def fill_step1(wb, clips, class_order):
    """Rebuild Step1 sheet rows so any number of clips fits; returns (total_cell, acc_cell)."""
    ws = wb["Step1_Individual_Clips"]
    S = lambda ref: copy(ws[ref]._style)
    st_in = {c: S(f"{c}6") for c in "ABCDF"}
    st_e = S("E6")
    st_sec, st_lbl = S("A57"), S("A58")
    st_val = {r: S(f"B{r}") for r in range(58, 63)}
    st_th = {c: S(f"{c}67") for c in "ABCD"}
    st_cls = {c: S(f"{c}68") for c in "ABCD"}
    st_note = S("A66")

    for row in range(5, max(ws.max_row, 100) + 1):
        for col in range(1, 8):
            _reset(ws.cell(row, col))

    ws["A2"] = ("Auto-filled by make_test_report.py: Clip Filename, True Label, Predicted Label, Confidence and Notes. "
                "'Correct?' and all summary stats compute in Excel. Predicted 'uncertain' = no class reached the confidence threshold.")

    n = max(len(clips), 1)
    last = 4 + n
    for i, c in enumerate(clips):
        r = 5 + i
        for col, key in zip("ABCDF", ("file", "true", "pred", "conf", "note")):
            cell = ws[f"{col}{r}"]
            cell._style = copy(st_in[col])
            cell.value = round(c[key], 4) if key == "conf" else c[key]
        ws[f"E{r}"]._style = copy(st_e)
        ws[f"E{r}"] = f'=IF(B{r}="","",IF(B{r}=C{r},"Yes","No"))'
    if not clips:
        for col in "ABCDF":
            ws[f"{col}5"]._style = copy(st_in[col])
        ws["E5"]._style = copy(st_e)

    sr = last + 3
    rng = lambda col: f"{col}5:{col}{last}"
    rows = [
        ("Total Clips Tested", f"=COUNTA({rng('B')})", 58),
        ("Correct", f'=COUNTIF({rng("E")},"Yes")', 59),
        ("Overall Accuracy", f'=IFERROR(B{sr + 2}/B{sr + 1},"")', 60),
        ("Avg. Confidence (correct predictions)", f'=IFERROR(AVERAGEIF({rng("E")},"Yes",{rng("D")}),"")', 61),
        ("Avg. Confidence (wrong predictions)", f'=IFERROR(AVERAGEIF({rng("E")},"No",{rng("D")}),"")', 62),
    ]
    ws[f"A{sr}"]._style = copy(st_sec)
    ws[f"A{sr}"] = "Overall Summary"
    for k, (lab, f, tr) in enumerate(rows, start=1):
        ws[f"A{sr + k}"]._style = copy(st_lbl)
        ws[f"A{sr + k}"] = lab
        ws[f"B{sr + k}"]._style = copy(st_val[tr])
        ws[f"B{sr + k}"] = f

    pr = sr + 8
    ws[f"A{pr}"]._style = copy(st_sec)
    ws[f"A{pr}"] = "Per-Class Breakdown"
    ws[f"A{pr + 1}"]._style = copy(st_note)
    ws[f"A{pr + 1}"] = "One row per class (True Label)."
    for col, t in zip("ABCD", ("Class Name", "Tested", "Correct", "Accuracy")):
        ws[f"{col}{pr + 2}"]._style = copy(st_th[col])
        ws[f"{col}{pr + 2}"] = t
    nrows = max(15, len(class_order))
    for k in range(nrows):
        r = pr + 3 + k
        for col in "ABCD":
            ws[f"{col}{r}"]._style = copy(st_cls[col])
        ws[f"A{r}"] = class_order[k] if k < len(class_order) else None
        ws[f"B{r}"] = f'=IF(A{r}="","",COUNTIF(B$5:B${last},A{r}))'
        ws[f"C{r}"] = f'=IF(A{r}="","",COUNTIFS(B$5:B${last},A{r},E$5:E${last},"Yes"))'
        ws[f"D{r}"] = f'=IF(A{r}="","",IFERROR(C{r}/B{r},""))'
    return f"B{sr + 1}", f"B{sr + 3}"


def fill_cycle_sheet(wb, proto, sheet_name, title, gt_segs, det_segs, matches, det_notes):
    ws = wb.copy_worksheet(proto)
    ws.title = sheet_name
    ws["A1"] = title
    S = lambda ref: copy(ws[ref]._style)
    # example rows -> normal input style
    for c in "ABCDE":
        ws[f"{c}8"]._style = copy(ws[f"{c}9"]._style)
    for c in "ABCDEF":
        ws[f"{c}52"]._style = copy(ws[f"{c}53"]._style)
    for ref in ("B8", "C8", "D8", "E8", "B52", "C52", "D52", "F52"):
        ws[ref].value = None
    for i, g in enumerate(gt_segs[:MAX_ROWS]):
        r = 8 + i
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = g["label"], g["start"], g["end"]
    for i, d in enumerate(det_segs[:MAX_ROWS]):
        r = 52 + i
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = d["label"], round(d["start"], 2), round(d["end"], 2)
        if det_notes.get(i):
            ws[f"F{r}"] = det_notes[i]
    for i, m in enumerate(matches[:MAX_ROWS]):
        ws[f"E{96 + i}"] = m
    return ws


# ----------------------------------------------------------------------------- main
def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--videomae_dir", required=True)
    p.add_argument("--individual_dir", default=None, help="folder of single-class clips (Step 1)")
    p.add_argument("--continuous_dir", default=None, help="folder of cycleN videos (Step 2)")
    p.add_argument("--gt", default="cycle20GT.ods")
    p.add_argument("--template", default="temporal_test_template.xlsx")
    p.add_argument("--window_len", type=float, default=2.0)
    p.add_argument("--stride", type=float, default=1.0)
    p.add_argument("--fps", type=float, default=15.0)
    p.add_argument("--num_frames", type=int, default=16)
    p.add_argument("--min_segment_duration", type=float, default=2.5)
    p.add_argument("--confidence_threshold", type=float, default=None)
    p.add_argument("--roi", default=None)
    p.add_argument("--runs_dir", default="runs")
    p.add_argument("--out_dir", default=None)
    p.add_argument("--out_name", default=None)
    p.add_argument("--clip_method", choices=["dominant_segment", "mean_prob"], default="dominant_segment",
                   help="how one label/confidence is derived for a single-class clip (Step 1)")
    p.add_argument("--min_overlap", type=float, default=1.0,
                   help="min seconds of overlap for a detected segment to be matched to a GT step (Step 2)")
    p.add_argument("--keep_uncertain", action="store_true",
                   help="list 'uncertain' segments in Section B too (they then count as false positives)")
    p.add_argument("--merge_step6", action="store_true",
                   help="merge step6 into step6.5 (default: they are separate tests)")
    p.add_argument("--tester", default=None)
    p.add_argument("--test_date", default=None)
    p.add_argument("--no_cache", action="store_true")
    return p.parse_args()


def load_backend():
    import types
    from extract_embeddings import extract_session_embeddings, resolve_roi
    from segment_utils import apply_min_duration_filter, labels_to_segments
    from pipeline_runner import compute_model_fingerprint, VIDEO_EXTENSIONS
    return types.SimpleNamespace(**locals())


def list_videos(folder, exts):
    return [os.path.join(folder, f) for f in sorted(os.listdir(folder)) if f.lower().endswith(tuple(exts))]


def main():
    a = parse_args()
    if not a.individual_dir and not a.continuous_dir:
        sys.exit("Give --individual_dir and/or --continuous_dir")
    rename = dict(MERGE_STEP6_RENAME) if a.merge_step6 else {}
    be = load_backend()

    import openpyxl
    fingerprint = be.compute_model_fingerprint(a.videomae_dir)
    roi = be.resolve_roi(a.videomae_dir, a.roi)
    with open(os.path.join(a.videomae_dir, "config.json")) as f:
        id2label = json.load(f)["id2label"]
    raw_names = [id2label[str(i)] for i in range(len(id2label))]
    uniq, M = collapse_matrix(raw_names, rename)
    print(f"Model fingerprint: {fingerprint}")
    print(f"Model classes ({len(raw_names)}): {raw_names}")
    print(f"Excel labels   ({len(uniq)}): {uniq}   (rename map: {rename or 'none'})\n")

    engine = Engine(a, be, raw_names, roi, fingerprint)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = a.out_dir or os.path.join(a.runs_dir, ts + "_report")
    os.makedirs(out_dir, exist_ok=True)

    wb = openpyxl.load_workbook(a.template)
    results = {"model_fingerprint": fingerprint, "individual": [], "continuous": []}
    notes_for_runinfo = []

    # ---------------- Step 1
    if a.individual_dir:
        print("=== STEP 1: individual clips ===")
        clips = []
        for vp in list_videos(a.individual_dir, be.VIDEO_EXTENSIONS):
            true = label_from_filename(vp, rename)
            if true is None:
                print(f"  [skip] cannot parse label from filename: {os.path.basename(vp)}")
                continue
            if true not in uniq:
                print(f"  [WARNING] true label '{true}' ({os.path.basename(vp)}) is not one of the model's labels {uniq}")
            wt, probs = engine.windows_and_probs(vp)
            pc = np.asarray(probs) @ M if len(wt) else np.zeros((0, len(uniq)))
            segs, _labels, conf = windows_to_segments(wt, pc, uniq, a, be)
            lab, c, note = clip_prediction(wt, pc, uniq, segs, conf, a)
            clips.append({"file": os.path.basename(vp), "true": true, "pred": lab, "conf": c, "note": note})
            print(f"  {os.path.basename(vp):34s} true={true:11s} pred={lab:11s} conf={c:.3f} {'' if lab == true else '  <-- WRONG'}")
        seen = list(dict.fromkeys(c["true"] for c in clips))
        class_order = [u for u in uniq if u in seen] + [s for s in seen if s not in uniq]
        tot_cell, acc_cell = fill_step1(wb, clips, class_order)
        results["individual"] = clips
        ncor = sum(c["pred"] == c["true"] for c in clips)
        print(f"\n  Step 1 accuracy: {ncor}/{len(clips)} = {ncor / max(len(clips), 1):.1%}")
        for cl in class_order:
            sub = [c for c in clips if c["true"] == cl]
            print(f"    {cl:12s} {sum(c['pred'] == cl for c in sub)}/{len(sub)}")
        print()
        wb["Summary_Dashboard"]["B8"] = f"=Step1_Individual_Clips!{tot_cell}"
        wb["Summary_Dashboard"]["B9"] = f"=Step1_Individual_Clips!{acc_cell}"
        n_unc = sum(c["pred"] == UNCERTAIN for c in clips)
        notes_for_runinfo.append(f"Step1: {len(clips)} clips, method={a.clip_method}, {n_unc} predicted 'uncertain'.")

    # ---------------- Step 2
    if a.continuous_dir:
        print("=== STEP 2: continuous videos ===")
        gt, gt_labels = load_gt(a.gt, rename)
        bad = sorted(set(l for segs in gt.values() for l in [s["label"] for s in segs]) - set(uniq))
        if bad:
            print(f"  [WARNING] GT labels not among model labels (will show as MISSED / WRONG): {bad}")
        proto = wb["Step2_Continuous_Video"]
        vids = list_videos(a.continuous_dir, be.VIDEO_EXTENSIONS)
        vids.sort(key=lambda p: int(re.search(r"(\d+)", os.path.basename(p)).group(1)))
        cyc_sheets, trunc = [], []
        for vp in vids:
            m = re.search(r"cycle[_\-]?(\d+)", os.path.basename(vp), re.I)
            if not m or int(m.group(1)) not in gt:
                print(f"  [skip] no GT row for {os.path.basename(vp)}")
                continue
            cyc = int(m.group(1))
            wt, probs = engine.windows_and_probs(vp)
            pc = np.asarray(probs) @ M
            segs, _l, _c = windows_to_segments(wt, pc, uniq, a, be)
            det_all = [{"label": s["label"], "start": s["start_sec"], "end": s["end_sec"]} for s in segs]
            det = det_all if a.keep_uncertain else [d for d in det_all if d["label"] != UNCERTAIN]
            n_unc = len(det_all) - len(det)
            g = gt[cyc]
            if len(det) > MAX_ROWS:
                trunc.append((cyc, len(det)))
            matches = match_gt_to_det(g, det[:MAX_ROWS], a.min_overlap)

            sheet = f"Step2_cycle{cyc:02d}"
            fill_cycle_sheet(wb, proto, sheet,
                             f"Step 2 -- Continuous Video: cycle {cyc} ({os.path.basename(vp)})", g, det, matches, {})
            cyc_sheets.append((cyc, sheet))

            ok = sum(1 for gg, mm in zip(g, matches) if mm and det[mm - 1]["label"] == gg["label"])
            miss = sum(1 for mm in matches if mm is None)
            wrong = len(g) - ok - miss
            used = {mm for mm in matches if mm}
            fp = sum(1 for i in range(min(len(det), MAX_ROWS)) if (i + 1) not in used)
            results["continuous"].append({"cycle": cyc, "video": vp, "gt": g, "detected": det, "matches": matches,
                                          "ok": ok, "missed": miss, "wrong_label": wrong, "false_positives": fp,
                                          "uncertain_segments_dropped": n_unc})
            print(f"  cycle {cyc:2d}: GT={len(g)} OK={ok} missed={miss} wrong={wrong} FP={fp} "
                  f"(detected {len(det)}, dropped {n_unc} uncertain)")
        del wb["Step2_Continuous_Video"]
        if trunc:
            print(f"  [WARNING] more than {MAX_ROWS} detected segments (template capacity), extras NOT in Excel: {trunc}")
        tot = {k: sum(r[k] for r in results["continuous"]) for k in ("ok", "missed", "wrong_label", "false_positives")}
        tot_gt = sum(len(r["gt"]) for r in results["continuous"])
        print(f"\n  Step 2 overall: {tot['ok']}/{tot_gt} = {tot['ok'] / max(tot_gt, 1):.1%} OK, "
              f"missed={tot['missed']}, wrong label={tot['wrong_label']}, false positives={tot['false_positives']}\n")
        build_dashboard_step2(wb, cyc_sheets)
        notes_for_runinfo.append(f"Step2: {len(cyc_sheets)} cycles, one sheet each; 'uncertain' segments "
                                 f"{'kept' if a.keep_uncertain else 'excluded'} from Section B; match needs >= {a.min_overlap}s overlap.")
        if trunc:
            notes_for_runinfo.append(f"TRUNCATED detections (>{MAX_ROWS}) in cycles {trunc}")

    # ---------------- Run Info
    ri = wb["Run Info"]
    run_name = os.path.basename(os.path.dirname(os.path.dirname(os.path.abspath(a.videomae_dir.rstrip('/')))))
    ri["B6"] = a.test_date or datetime.now().strftime("%Y-%m-%d")
    if a.tester:
        ri["B7"] = a.tester
    ri["B8"] = run_name
    ri["B9"] = fingerprint
    ri["B10"] = ",".join(str(int(v)) if float(v).is_integer() else str(v) for v in roi) if roi else "none"
    ri["B11"], ri["B12"], ri["B13"] = a.window_len, a.stride, a.num_frames
    ri["B14"] = a.confidence_threshold
    ri["B15"] = a.min_segment_duration
    ri["B16"] = (f"fps={a.fps}. Label rename: {rename or 'none'}. " + " ".join(notes_for_runinfo))

    wb.calculation.fullCalcOnLoad = True
    out_name = a.out_name or f"temporal_test_report_{run_name}.xlsx"
    out_path = os.path.join(out_dir, out_name)
    wb.save(out_path)
    with open(os.path.join(out_dir, "results.json"), "w") as f:
        json.dump(results, f, indent=2, default=str)
    print(f"Saved: {out_path}\nRaw results: {os.path.join(out_dir, 'results.json')}")


def build_dashboard_step2(wb, cyc_sheets):
    ws = wb["Summary_Dashboard"]
    S = lambda ref: copy(ws[ref]._style)
    st_l, st_v, st_pct = S("A14"), S("B14"), S("B13")
    hdr = copy(wb["Step1_Individual_Clips"]["A4"]._style)
    body = copy(wb["Step1_Individual_Clips"]["A58"]._style)

    sh = [s for _, s in cyc_sheets]
    plus = lambda cell: "+".join(f"{s}!{cell}" for s in sh)
    ws["A11"] = f"STEP 2 -- Continuous Video ({len(sh)} cycles combined)"
    ws["B12"] = "=" + plus("B139")
    ws["B13"] = f'=IFERROR(({plus("B140")})/B12,"")'
    ws["B14"] = "=" + plus("B144")
    extra = [
        (15, "Missed Steps", "=" + plus("B141"), st_v),
        (16, "Wrong Label Steps", "=" + plus("B142"), st_v),
        (17, "Avg. |Start Offset| for OK matches (s)",
         "=IFERROR((" + "+".join(f'SUMIF({s}!J96:J135,"OK",{s}!I96:I135)' for s in sh) + f")/({plus('B140')}),\"\")", st_v),
    ]
    for r, lab, f, st in extra:
        ws[f"A{r}"]._style, ws[f"A{r}"] = copy(st_l), lab
        ws[f"B{r}"]._style, ws[f"B{r}"] = copy(st), f
    ws["B17"].number_format = "0.00"

    r0 = 20
    ws[f"A{r0 - 1}"] = "Per-cycle results"
    ws[f"A{r0 - 1}"].font = copy(wb["Step1_Individual_Clips"]["A57"].font)
    for col, t in zip("ABCDEFG", ("Cycle", "GT Steps", "OK", "Missed", "Wrong Label", "False Positives", "Accuracy")):
        ws[f"{col}{r0}"]._style, ws[f"{col}{r0}"] = copy(hdr), t
    for i, (cyc, s) in enumerate(cyc_sheets, start=1):
        r = r0 + i
        vals = [f"cycle {cyc}", f"={s}!B139", f"={s}!B140", f"={s}!B141", f"={s}!B142", f"={s}!B144",
                f'=IFERROR({s}!B140/{s}!B139,"")']
        for col, v in zip("ABCDEFG", vals):
            ws[f"{col}{r}"]._style, ws[f"{col}{r}"] = copy(body), v
        ws[f"G{r}"].number_format = "0.0%"
    for col in "DEFG":
        ws.column_dimensions[col].width = 16


if __name__ == "__main__":
    main()
